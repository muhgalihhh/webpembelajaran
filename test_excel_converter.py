#!/usr/bin/env python3
"""
Test script untuk Excel Preview & PDF Converter
"""

import os
import sys
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

def create_test_excel():
    """Membuat file Excel test dengan data sample"""
    
    # Buat workbook baru
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Sample"
    
    # Buat sheet kedua
    ws2 = wb.create_sheet("Laporan")
    
    # Data untuk sheet pertama
    headers = ["Nama", "Usia", "Kota", "Pekerjaan", "Gaji"]
    data = [
        ["John Doe", 30, "Jakarta", "Programmer", 5000000],
        ["Jane Smith", 25, "Bandung", "Designer", 4000000],
        ["Bob Johnson", 35, "Surabaya", "Manager", 8000000],
        ["Alice Brown", 28, "Medan", "Analyst", 6000000],
        ["Charlie Wilson", 32, "Semarang", "Developer", 5500000],
    ]
    
    # Tambahkan header
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    # Tambahkan data
    for row, row_data in enumerate(data, 2):
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
    
    # Data untuk sheet kedua
    ws2['A1'] = "LAPORAN KEUANGAN"
    ws2['A1'].font = Font(bold=True, size=16)
    ws2['A1'].alignment = Alignment(horizontal="center")
    
    ws2['A3'] = "Periode: Januari 2024"
    ws2['A3'].font = Font(bold=True)
    
    ws2['A5'] = "Pendapatan"
    ws2['B5'] = 15000000
    ws2['A6'] = "Pengeluaran"
    ws2['B6'] = 8000000
    ws2['A7'] = "Laba"
    ws2['B7'] = "=B5-B6"
    
    # Format angka
    for cell in [ws2['B5'], ws2['B6'], ws2['B7']]:
        cell.number_format = '#,##0'
    
    # Buat file temporary
    temp_dir = tempfile.gettempdir()
    test_file = os.path.join(temp_dir, "test_excel_converter.xlsx")
    
    wb.save(test_file)
    print(f"File test Excel telah dibuat: {test_file}")
    
    return test_file

def main():
    """Fungsi utama untuk testing"""
    print("=== Excel Preview & PDF Converter Test ===")
    print()
    
    try:
        # Buat file test Excel
        test_file = create_test_excel()
        
        print("File test berhasil dibuat!")
        print(f"Path: {test_file}")
        print()
        print("Sekarang Anda dapat:")
        print("1. Jalankan aplikasi GUI: python excel_preview_converter_simple.py")
        print("2. Pilih file test yang baru dibuat")
        print("3. Test fitur preview dan konversi ke PDF")
        print()
        print("Atau gunakan script konversi langsung:")
        print(f"python -c \"import win32com.client; excel = win32com.client.Dispatch('Excel.Application'); excel.Visible = False; wb = excel.Workbooks.Open(r'{test_file}'); wb.ExportAsFixedFormat(0, r'{test_file.replace('.xlsx', '.pdf')}'); wb.Close(False); excel.Quit(); print('PDF berhasil dibuat!')\"")
        
    except ImportError:
        print("Error: openpyxl tidak terinstall")
        print("Install dengan: pip install openpyxl")
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    main()